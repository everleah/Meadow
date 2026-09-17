"""`outbound_export` job handler: `json_export` connector (MEF export, Phase 6).

The handler receives `parameters.filter` (`type` and/or `status`), fetches the
matching objects from the core, builds a Meadow Exchange Format (MEF) document,
and stores it in `job_run.result`.

Quantified values (stored as canonical-SI numbers) are re-emitted as
`{"value": …, "unit": …}` in the attribute's display unit via the core's
`GET /units/convert` endpoint.
"""

from __future__ import annotations

import base64
import csv as _csv
import io
import json as _json
import logging
import os
from collections.abc import Iterable, Mapping
from datetime import UTC, datetime
from typing import Any, NamedTuple

import httpx
from openpyxl import Workbook
from openpyxl.styles import Font

from meadow.urls import is_valid_url
from meadow.worker import bmecat
from meadow.worker.http import loopback_host
from meadow.worker.jobs import JobHandler, JobMessage
from meadow.worker.mail import SmtpConfig, send_with_attachment
from meadow.worker.storage import S3Config, archive_export

logger = logging.getLogger(__name__)

#: Default cap for exported objects per run.
DEFAULT_EXPORT_LIMIT = 1000

#: Page size of the core's list endpoints (`limit` is capped there).
OBJECT_PAGE_SIZE = 1000

#: Object ids per `POST /references/query` call (core limit).
REFERENCE_QUERY_SIZE = 500

#: Rows per `POST /references/query` page (core default is 5000).
REFERENCE_QUERY_PAGE = 5000

#: Composition relation definition (record key in the core, `depends_on` as name).
COMPOSITION_DEFINITION = "rel_depends_on"

#: Object type of the universal root (a system object, never part of a document).
ROOT_TYPE = "root"

#: Record key of the universal root (exists in every instance).
ROOT_KEY = "object_root"

#: Upper bound of composition levels the ancestor walk follows (a cycle guard).
MAX_ANCESTOR_ROUNDS = 50

#: Font of a hyperlinked export cell (blue + underlined, like a browser link).
URL_CELL_FONT = Font(color="FF0563C1", underline="single")


class AttributeFlags(NamedTuple):
    """Attribute types that change how an attribute is exported."""

    #: Attributes holding one value per language (column per locale).
    multilingual: set[str]
    #: Attributes stored as canonical-SI numbers with a unit.
    quantified: set[str]
    #: Attributes of type `url` (XLSX cells become clickable hyperlinks).
    urls: set[str]


def _now_iso_utc() -> str:
    """Returns the current time as an ISO-8601 UTC string."""
    return datetime.now(UTC).isoformat()


def _is_number(value: Any) -> bool:
    """Whether `value` is a JSON number (excluding booleans)."""
    return isinstance(value, (int, float)) and not isinstance(value, bool)


async def _load_unit_map(
    client: httpx.AsyncClient,
    core_url: str,
) -> dict[str, tuple[str, str]]:
    """Loads quantified attribute definitions: `name -> (unit_id, display_unit_id)`."""
    response = await client.get(
        f"{core_url}/attribute-definitions",
        params={"limit": 1000},
    )
    response.raise_for_status()
    unit_map: dict[str, tuple[str, str]] = {}
    for definition in response.json():
        name = definition.get("name")
        unit_id = definition.get("unit_id")
        if not name or not unit_id:
            continue
        display_unit_id = definition.get("display_unit_id") or unit_id
        unit_map[str(name)] = (str(unit_id), str(display_unit_id))
    return unit_map


async def _emit_quantified(
    client: httpx.AsyncClient,
    core_url: str,
    attributes: dict[str, Any],
    unit_map: dict[str, tuple[str, str]],
) -> None:
    """Re-emits canonical-SI numbers as `{"value":…,"unit":…}` objects."""
    for name, (unit_id, display_unit_id) in unit_map.items():
        value = attributes.get(name)
        if not _is_number(value):
            continue
        response = await client.get(
            f"{core_url}/units/convert",
            params={"value": value, "from": unit_id, "to": display_unit_id},
        )
        response.raise_for_status()
        converted = response.json()
        attributes[name] = {"value": converted["value"], "unit": converted["unit"]}


# --- Structured export (UI wizard, Phase 9) --------------------------------

_BASE_FIELDS = ("id", "type", "status")


def _strip_object_id(value: Any) -> str:
    """Strips an `object:` prefix, keeping the bare record key."""
    text = str(value)
    return text.split(":", 1)[1] if ":" in text else text


def _chunks(items: list[Any], size: int) -> Iterable[list[Any]]:
    """Yields `items` in chunks of at most `size` entries."""
    for start in range(0, len(items), size):
        yield items[start : start + size]


async def _page_get(
    active: httpx.AsyncClient,
    url: str,
    params: Mapping[str, Any] | None = None,
) -> list[Any]:
    """Reads a `limit`/`start` paged list endpoint **completely**.

    Every list endpoint of the core caps `limit` at [`OBJECT_PAGE_SIZE`], so an
    export that asked for one page silently stopped at 1 000 rows — a 2 000-product
    catalogue lost half of its objects, which is why a "complete" export was not
    one. The `id`s of a page are tracked as a safety net: if a page brings nothing
    new (an unstable order), the loop stops instead of running forever.
    """
    rows: list[Any] = []
    seen: set[str] = set()
    start = 0
    while True:
        response = await active.get(
            url,
            params={**(params or {}), "limit": OBJECT_PAGE_SIZE, "start": start},
        )
        response.raise_for_status()
        payload = response.json()
        page = payload if isinstance(payload, list) else []
        fresh = 0
        for row in page:
            object_id = str(row.get("id")) if isinstance(row, dict) and row.get("id") else ""
            if not object_id:
                fresh += 1
            elif object_id not in seen:
                seen.add(object_id)
                fresh += 1
        rows.extend(page)
        if len(page) < OBJECT_PAGE_SIZE or fresh == 0:
            return rows
        start += OBJECT_PAGE_SIZE


async def _page_objects(
    active: httpx.AsyncClient,
    core_url: str,
    params: Mapping[str, Any],
) -> list[dict[str, Any]]:
    """Every object of a `GET /objects` query (see [`_page_get`] for the paging)."""
    return [
        row
        for row in await _page_get(active, f"{core_url}/objects", params)
        if isinstance(row, dict)
    ]


async def _load_objects_by_ids(
    active: httpx.AsyncClient,
    core_url: str,
    ids: Iterable[Any],
    *,
    extra: Mapping[str, Any] | None = None,
) -> list[dict[str, Any]]:
    """Reads objects by id, chunked by the core's `?ids=` page size."""
    objects: list[dict[str, Any]] = []
    for chunk in _chunks([str(item) for item in ids], OBJECT_PAGE_SIZE):
        objects.extend(
            await _page_objects(active, core_url, {"ids": ",".join(chunk), **(extra or {})})
        )
    return objects


async def _resolve_objects(
    active: httpx.AsyncClient,
    core_url: str,
    source: dict[str, Any],
    *,
    resolve_inheritance: bool = False,
    with_releases: bool = False,
) -> list[dict[str, Any]]:
    """Resolves the export source to a list of full object records.

    Every source kind is read **completely** (`_page_get`): the list endpoints cap
    `limit` at 1000, so a type with more objects used to be truncated silently.
    `kind='all'` is the complete export — every live object except the universal
    root, which is a system object present in every instance.

    With ``resolve_inheritance`` the objects are requested with
    `?resolve_inheritance=true`, so attributes maintained on an ancestor are
    exported with the child's effective value. With ``with_releases`` they are
    requested with `?releases=true`: the channel releases are part of a MEF
    document (a restore without them has released nothing). `source.kind='channel'`
    exports the objects **released for that channel** (`?released_for=`, optionally
    with `source.release_status`).
    """
    extra: dict[str, str] = {"resolve_inheritance": "true"} if resolve_inheritance else {}
    if with_releases:
        extra["releases"] = "true"
    kind = source.get("kind")
    objects: list[dict[str, Any]]
    if kind == "all":
        objects = await _page_objects(active, core_url, dict(extra))
    elif kind == "type":
        object_type = source.get("type")
        if not object_type:
            raise RuntimeError("`source.kind='type'` requires `type`")
        objects = await _page_objects(active, core_url, {"type": object_type, **extra})
    elif kind == "category":
        category = source.get("category")
        if not category:
            raise RuntimeError("`source.kind='category'` requires `category`")
        objects = await _page_objects(active, core_url, {"category": category, **extra})
    elif kind == "channel":
        channel = source.get("channel")
        if not channel:
            raise RuntimeError("`source.kind='channel'` requires `channel`")
        params: dict[str, Any] = {"released_for": channel, **extra}
        release_status = source.get("release_status")
        if release_status:
            params["release_status"] = str(release_status)
        objects = await _page_objects(active, core_url, params)
    else:
        ids: list[str] = []
        if kind == "collection":
            collection_id = source.get("id")
            if not collection_id:
                raise RuntimeError("`source.kind='collection'` requires `id`")
            rows = await _page_get(active, f"{core_url}/collections/{collection_id}/members")
            ids = [_strip_object_id(item.get("id")) for item in rows]
        elif kind == "saved_search":
            search_id = source.get("id")
            if not search_id:
                raise RuntimeError("`source.kind='saved_search'` requires `id`")
            resp = await active.post(f"{core_url}/saved-searches/{search_id}/execute")
            resp.raise_for_status()
            ids = [_strip_object_id(item) for item in (resp.json().get("ids") or [])]
        elif kind == "selection":
            ids = [_strip_object_id(item) for item in (source.get("object_ids") or [])]
        else:
            raise RuntimeError(f"unsupported `source.kind` {kind!r}")
        objects = await _load_objects_by_ids(active, core_url, ids, extra=extra) if ids else []
    # The universal root is a system object that exists in every instance — it is
    # never part of an export (and never a parent a document has to carry).
    return [obj for obj in objects if str(obj.get("type") or "") != ROOT_TYPE]


async def _load_languages(client: httpx.AsyncClient, core_url: str) -> dict[str, list[str]]:
    """Returns `code -> fallback_order` from the language tree."""
    resp = await client.get(f"{core_url}/languages")
    resp.raise_for_status()
    result: dict[str, list[str]] = {}
    for lang in resp.json():
        code = lang.get("code")
        if code:
            result[str(code)] = [str(c) for c in (lang.get("fallback_order") or [])]
    return result


async def _load_attribute_flags(client: httpx.AsyncClient, core_url: str) -> AttributeFlags:
    """Returns the attribute-type flags the export needs in one catalog call."""
    resp = await client.get(f"{core_url}/attribute-definitions", params={"limit": 1000})
    resp.raise_for_status()
    multilingual: set[str] = set()
    quantified: set[str] = set()
    urls: set[str] = set()
    for definition in resp.json():
        name = definition.get("name")
        if not name:
            continue
        if definition.get("multilingual"):
            multilingual.add(str(name))
        if definition.get("unit_id"):
            quantified.add(str(name))
        if definition.get("type") == "url":
            urls.add(str(name))
    return AttributeFlags(multilingual=multilingual, quantified=quantified, urls=urls)


def _pick(value: Any, fallback: dict[str, list[str]], locale: str) -> Any:
    """Resolves a multilingual dict via the locale fallback chain.

    A dict is treated as a language map only when at least one of its keys is a
    known language code; structured dicts (e.g. ``{"value":…,"unit":…}``) are
    returned unchanged.
    """
    if not isinstance(value, dict):
        return value
    if not any(key in fallback or key == locale for key in value):
        return value
    chain = [locale, *fallback.get(locale, []), "en"]
    for code in chain:
        if code in value and value[code] is not None:
            return value[code]
    return ""


def _build_columns(fields: list[str], locales: list[str], multilingual: set[str]) -> list[str]:
    """Builds the column list, expanding multilingual attrs per locale."""
    columns: list[str] = []
    for field in fields:
        if field in _BASE_FIELDS or field not in multilingual:
            columns.append(field)
        else:
            for locale in locales:
                columns.append(f"{field}_{locale}")
    if not columns:
        columns = list(_BASE_FIELDS)
    return columns


def _build_rows(
    objects: list[dict[str, Any]],
    columns: list[str],
    locales: list[str],
    fallback: dict[str, list[str]],
    multilingual: set[str],
) -> list[dict[str, Any]]:
    """Builds a row per object, resolving multilingual columns per locale."""
    rows: list[dict[str, Any]] = []
    for obj in objects:
        attrs = obj.get("attributes") or {}
        attributes = attrs if isinstance(attrs, dict) else {}
        row: dict[str, Any] = {}
        for column in columns:
            if column == "id":
                row["id"] = _strip_object_id(obj.get("id"))
            elif column in ("type", "status"):
                row[column] = obj.get(column) or ""
            elif "_" in column and column.rsplit("_", 1)[0] in multilingual:
                # multilingual column `<attr>_<locale>`
                attr_name, locale = column.rsplit("_", 1)
                row[column] = _pick(attributes.get(attr_name), fallback, locale)
            else:
                # plain attribute column: resolve dicts with the first locale
                row[column] = _pick(
                    attributes.get(column, ""), fallback, locales[0] if locales else "de"
                )
        rows.append(row)
    return rows


def _serialize_csv(
    columns: list[str], rows: list[dict[str, Any]], separator: str, header: bool
) -> str:
    """Serializes rows to CSV (UTF-8)."""
    buffer = io.StringIO()
    writer = _csv.writer(buffer, delimiter=separator or ",")
    if header:
        writer.writerow(columns)
    for row in rows:
        writer.writerow([_csv_cell(row.get(col, "")) for col in columns])
    return buffer.getvalue()


def _csv_cell(value: Any) -> str:
    """Converts a cell value to a CSV string (JSON for structured values)."""
    if isinstance(value, (dict, list)):
        return _json.dumps(value, ensure_ascii=False)
    if isinstance(value, bool):
        return "true" if value else "false"
    return "" if value is None else str(value)


def _serialize_xlsx(
    columns: list[str], rows: list[dict[str, Any]], url_columns: set[str] | None = None
) -> bytes:
    """Serializes rows to an XLSX workbook and returns the bytes.

    Cells of `url` attribute columns carry a **hyperlink**, so the exported file
    links straight to the website/reference (a `text` attribute holding a URL
    stays plain text — the data type is the contract). The value itself is the
    URL either way, so a CSV/JSON export keeps working unchanged.
    """
    links = url_columns or set()
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:
        sheet = workbook.create_sheet(title="Export")
    else:
        sheet.title = "Export"
    sheet.append(columns)
    for offset, row in enumerate(rows, start=2):
        for index, column in enumerate(columns, start=1):
            text = _csv_cell(row.get(column, ""))
            cell = sheet.cell(row=offset, column=index, value=text)
            if column in links and is_valid_url(text):
                cell.hyperlink = text
                cell.font = URL_CELL_FONT
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _url_columns(columns: list[str], urls: set[str]) -> set[str]:
    """Columns whose values are URLs (multilingual columns are `<name>_<locale>`)."""
    result: set[str] = set()
    for column in columns:
        if column in urls or ("_" in column and column.rsplit("_", 1)[0] in urls):
            result.add(column)
    return result


def _attachment_from_result(result: dict[str, Any]) -> tuple[str, bytes, str]:
    """Returns ``(filename, content, content_type)`` for an export result."""
    format_ = str(result.get("format") or "csv")
    if format_ == "xlsx":
        content = base64.b64decode(result.get("xlsx_base64") or "")
        return (
            "meadow-export.xlsx",
            content,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    if format_ == "json":
        body = _json.dumps(result.get("document") or {}, ensure_ascii=False, indent=2).encode(
            "utf-8"
        )
        return ("meadow-export.json", body, "application/json")
    if format_ == "bmecat":
        body = str(result.get("xml") or "").encode("utf-8")
        return ("meadow-export-bmecat.xml", body, "application/xml")
    return ("meadow-export.csv", str(result.get("csv") or "").encode("utf-8"), "text/csv")


async def _maybe_email(
    result: dict[str, Any],
    options: dict[str, Any],
    filename: str,
    content: bytes,
    content_type: str,
) -> None:
    """E-mails the export as an attachment when recipients and SMTP are configured.

    Best-effort: a transport failure (or missing SMTP config) is recorded under
    ``result["email"]`` but does not fail the export.
    """
    recipients = [str(r).strip() for r in (options.get("email_to") or []) if str(r).strip()]
    if not recipients:
        return
    config = SmtpConfig.from_env()
    if config is None:
        result["email"] = {"sent": 0, "error": "SMTP is not configured (SMTP_HOST unset)"}
        logger.warning("export e-mail requested but SMTP is not configured")
        return
    subject = str(options.get("email_subject") or "Meadow export")
    body = (
        f"Your Meadow export is attached ({result.get('exported', 0)} objects, "
        f"{str(result.get('format') or 'csv').upper()})."
    )
    try:
        await send_with_attachment(
            config, recipients, subject, body, filename, content, content_type
        )
    except Exception as exc:  # noqa: BLE001 - best-effort notification
        result["email"] = {"sent": 0, "error": str(exc)}
        logger.warning("export e-mail delivery failed: %s", exc)
        return
    result["email"] = {"sent": len(recipients)}


async def _maybe_archive(
    result: dict[str, Any], filename: str, content: bytes, content_type: str
) -> None:
    """Archives the export in object storage and drops large inline payloads.

    Best-effort: an upload failure is recorded under ``result["archive"]`` and
    the inline result is kept (the export itself still succeeds).
    """
    config = S3Config.from_env()
    if config is None:
        return
    stamp = datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")
    key = f"exports/{stamp}-{filename}"
    try:
        result["archive"] = await archive_export(config, key, content, content_type, filename)
    except Exception as exc:  # noqa: BLE001 - best-effort archive
        result["archive"] = {"error": str(exc)}
        logger.warning("export archive upload failed: %s", exc)
        return
    # Keep `job_run.result` small: drop the inline copy once it is archived and
    # exceeds the inline threshold (the download then uses the presigned URL).
    threshold = int(os.environ.get("EXPORT_INLINE_MAX_BYTES", "1000000"))
    if len(content) > threshold:
        for field in ("csv", "xlsx_base64", "document", "xml"):
            result.pop(field, None)


async def _finalize(result: dict[str, Any], options: dict[str, Any]) -> None:
    """Archives and/or e-mails a finished export (both best-effort)."""
    filename, content, content_type = _attachment_from_result(result)
    await _maybe_archive(result, filename, content, content_type)
    # The e-mail must use the pre-archive bytes (inline may have been dropped).
    await _maybe_email(result, options, filename, content, content_type)


async def _structured_export(
    active: httpx.AsyncClient,
    core_url: str,
    parameters: dict[str, Any],
) -> dict[str, Any]:
    """Runs a structured export and returns the serializable job result."""
    source = parameters.get("source") or {}
    format_ = str(parameters.get("format") or "json")
    if format_ == "bmecat":
        # The BMEcat mapping is fixed by the standard (class features, header
        # metadata) — no field selection, no CSV options.
        return await _bmecat_export(active, core_url, parameters)
    fields = [str(f) for f in (parameters.get("fields") or [])]
    if not fields:
        fields = list(_BASE_FIELDS)
    locales = [str(locale) for locale in (parameters.get("locales") or [])]
    if not locales:
        locales = ["de"]
    options = parameters.get("options") or {}
    header = bool(options.get("header", True))
    separator = str(options.get("separator") or ",")
    # "Also export inherited values": fills attributes maintained on an ancestor.
    resolve_inheritance = bool(options.get("resolve_inheritance"))

    flags = await _load_attribute_flags(active, core_url)
    multilingual = flags.multilingual
    fallback = await _load_languages(active, core_url)

    # Resolve objects (synchronously for the requested source).
    objects = await _resolve_objects(
        active,
        core_url,
        source,
        resolve_inheritance=resolve_inheritance,
        with_releases=format_ == "json",
    )
    if format_ == "json":
        unit_map = await _load_unit_map(active, core_url)
        # Relations (+ the composition ancestors of a partial source) are part of
        # the document: without them a fresh system gets a flat tree, and the
        # import refuses an object whose parent the document does not contain.
        objects, references, counters = await _collect_mef_data(
            active,
            core_url,
            objects,
            include_ancestors=bool(options.get("include_ancestors", True)),
            with_releases=True,
        )
        mef_objects = await _mef_objects(active, core_url, objects, unit_map)
        result: dict[str, Any] = {
            "format": "json",
            "exported": len(mef_objects),
            "columns": fields,
            "document": _mef_document(mef_objects, references),
            **counters,
        }
    else:
        columns = _build_columns(fields, locales, multilingual)  # type: ignore
        rows = _build_rows(objects, columns, locales, fallback, multilingual)
        if format_ == "xlsx":
            data = _serialize_xlsx(columns, rows, _url_columns(columns, flags.urls))
            result = {
                "format": "xlsx",
                "exported": len(objects),
                "columns": columns,
                "xlsx_base64": base64.b64encode(data).decode("ascii"),
            }
        else:
            csv_text = _serialize_csv(columns, rows, separator, header)
            result = {
                "format": "csv",
                "exported": len(objects),
                "columns": columns,
                "csv": csv_text,
            }

    await _finalize(result, options)
    return result


async def _bmecat_export(
    active: httpx.AsyncClient,
    core_url: str,
    parameters: dict[str, Any],
) -> dict[str, Any]:
    """Builds a BMEcat/ETIM catalogue document for the resolved source.

    Only objects of a product type with a supplier PID **and** at least one class
    become a `PRODUCT`; everything else is counted in `report["skipped"]` (the
    BMEcat profile identifies a product through its class, and `SUPPLIER_PID` is
    limited to 32 characters). The document is validated against the ETIM XSD when
    the schema and `lxml` are available — a missing schema is reported, not fatal
    (`plan/detail/bmecat-export.md` §8 V1).

    A `PRODUCT_FEATURES` block names its class by the **business key** (`code`,
    the ETIM class `EC…`) — an assigned taxonomy node without a code (a plain shop
    category) cannot be referenced in BMEcat and is therefore no block; those
    nodes are listed in `report["nodes_without_code"]` so a product that dropped
    out of the file can be explained.
    """
    profile = parameters.get("bmecat") or {}
    options = parameters.get("options") or {}
    definitions = {
        str(definition.get("name")): definition
        for definition in await _load_definitions(active, core_url)
    }
    # Which reference feature system does a class block name? The **installed**
    # standard decides (`GET /classification-standards`, `current` per code); the
    # profile's `etim_version` is only the fallback of an unregistered class.
    standards = await _load_standards(active, core_url)
    objects = await _resolve_objects(
        active,
        core_url,
        parameters.get("source") or {},
        resolve_inheritance=bool(options.get("resolve_inheritance")),
    )

    #: Node key -> (business key `code`, object type) of the assigned classes and
    #: of every node a feature comes from. One lookup per node, cached.
    class_nodes: dict[str, tuple[str, str]] = {}

    async def class_node(key: str) -> tuple[str, str]:
        """The `code` and the object type of a taxonomy node — `("", "")` unknown."""
        if key not in class_nodes:
            response = await active.get(f"{core_url}/objects/object:{key}")
            code, node_type = "", ""
            if response.status_code == 200:
                body = response.json()
                attributes = body.get("attributes") or {}
                code = str(attributes.get("code") or "").strip()
                node_type = str(body.get("type") or "").strip()
            class_nodes[key] = (code, node_type)
        return class_nodes[key]

    products: list[dict[str, Any]] = []
    without_code: set[str] = set()
    for obj in objects:
        attributes = dict(obj.get("attributes") or {})
        resolved = await _load_class_features(active, core_url, str(obj.get("id") or ""))
        keys = [str(key) for key in resolved.get("classes") or []]
        # One block per assigned class **with a business key**; the block names the
        # reference feature system of the standard installed for that node type.
        blocks: list[dict[str, str]] = []
        for key in keys:
            code, node_type = await class_node(key)
            if not code:
                without_code.add(key)
                continue
            blocks.append({"code": code, "system": standards.get(node_type, "")})
        # The features carry the **node keys** of the classes that define them; the
        # export identifies a class by its business key (`code`), so the membership
        # test compares codes (a key without a code — an ancestor of another
        # dimension — stays as-is and counts as inherited).
        features: list[dict[str, Any]] = []
        for feature in resolved.get("attributes") or []:
            entry = dict(feature)
            nodes: list[str] = []
            for node in feature.get("nodes") or []:
                code, _ = await class_node(str(node))
                nodes.append(code or str(node))
            entry["nodes"] = nodes
            features.append(entry)
        products.append(
            {
                "id": obj.get("id"),
                "type": obj.get("type"),
                "attributes": attributes,
                "pid": bmecat.pid_of(attributes, profile),
                # A variant references its base product when the base has a PID.
                "base_pid": await _base_pid(active, core_url, obj, profile),
                "classes": blocks,
                "features": features,
            }
        )

    document, report = bmecat.build_document(products, definitions, profile)
    report["nodes_without_code"] = sorted(without_code)
    report["standards"] = standards
    # The document's own version decides the schema to check it against (2005 is an
    # XSD, 1.2 ships a DTD and an XSD — `BMECAT_XSD` overrides both).
    report["validation"] = bmecat.validate(
        document,
        bmecat.schema_path(str(report.get("version") or ""), str(report.get("variant") or "")),
    )
    logger.info(
        "bmecat export wrote %d products (%d classes, %d features, validation=%s)",
        report["products"],
        report["classes"],
        report["features"],
        report["validation"].get("valid", report["validation"].get("validated")),
    )
    return {
        "format": "bmecat",
        "exported": report["products"],
        "catalog": str(profile.get("catalog_id") or "MEADOW"),
        "xml": document,
        "report": report,
    }


async def _load_definitions(client: httpx.AsyncClient, core_url: str) -> list[dict[str, Any]]:
    """Attribute definitions (type/unit/label) — the mapping needs their types."""
    response = await client.get(f"{core_url}/attribute-definitions", params={"limit": 1000})
    response.raise_for_status()
    return response.json()


async def _load_class_features(
    client: httpx.AsyncClient, core_url: str, object_id: str
) -> dict[str, Any]:
    """The resolved class features of an object (empty when it has no class)."""
    if not object_id:
        return {}
    response = await client.get(f"{core_url}/objects/{object_id}/category-attributes")
    if response.status_code != 200:
        return {}
    return response.json()


async def _base_pid(
    client: httpx.AsyncClient,
    core_url: str,
    obj: Mapping[str, Any],
    profile: Mapping[str, Any],
) -> str:
    """The `SUPPLIER_PID` of the base product of a **variant** (else `""`).

    BMEcat links a variant to its base product with
    `PRODUCT_REFERENCE type="base_product"`, and the reference target is the base
    product's PID (not its object id). The base product is the **composition
    parent** (`rel_depends_on`), which the core exposes through the object's
    references — there is no `GET /objects/{id}/parent` (only the re-parent PATCH).
    Only variants are looked up, so a catalogue pays two extra requests per
    variant, not per product.
    """
    if str(obj.get("type") or "") != "variant":
        return ""
    object_id = str(obj.get("id") or "")
    if not object_id:
        return ""
    response = await client.get(f"{core_url}/objects/{object_id}/references")
    if response.status_code != 200:
        return ""
    rows = response.json()
    parent = parent_reference_id(rows if isinstance(rows, list) else [])
    if not parent:
        return ""
    response = await client.get(f"{core_url}/objects/{parent}")
    if response.status_code != 200:
        return ""
    parent_object = response.json()
    if not isinstance(parent_object, dict):
        return ""
    return bmecat.pid_of(dict(parent_object.get("attributes") or {}), profile)


def parent_reference_id(
    references: Iterable[Mapping[str, Any]],
    *,
    composition: str = "rel_depends_on",
    root: str = "object:object_root",
) -> str:
    """The composition parent id of an object from its reference rows (`""` = root).

    The rows come from `GET /objects/{id}/references` (incoming *and* outgoing):
    the parent is the `out` of the composition edge, unless that is the object
    root (an object without a real base product). Pure, so it is unit-tested.
    """
    for row in references:
        definition = str(row.get("definition_id") or "")
        if definition.rsplit(":", 1)[-1] != composition:
            continue
        out = row.get("out")
        if not isinstance(out, str) or not out or out == root:
            continue
        return out
    return ""


async def _load_standards(client: httpx.AsyncClient, core_url: str) -> dict[str, str]:
    """`node_type` -> reference feature system of the **installed** standards.

    The registry (`GET /classification-standards`) knows which standard in which
    version was loaded for a node type; only the `current` version per standard
    counts (the newest — installing an update replaces the feature set by business
    key). A core without a registered standard, or an older core, yields an empty
    map: the profile's `etim_version` is then the fallback.
    """
    response = await client.get(f"{core_url}/classification-standards")
    if response.status_code != 200:
        return {}
    payload = response.json()
    rows = payload if isinstance(payload, list) else payload.get("items") or []
    systems: dict[str, str] = {}
    for row in rows:
        if not isinstance(row, dict) or not row.get("current"):
            continue
        node_type = str(row.get("node_type") or "").strip()
        name = str(row.get("feature_system_name") or "").strip()
        if node_type and name:
            systems[node_type] = name
    return systems


def _record_key(value: Any) -> str:
    """The bare record key of a `table:<key>` id (a bare key passes through)."""
    text = str(value or "")
    return text.split(":", 1)[1] if ":" in text else text


async def _load_relation_names(
    active: httpx.AsyncClient,
    core_url: str,
) -> dict[str, str]:
    """`relation_definition` record key -> MEF reference `type` (the definition name).

    A MEF reference carries the relation-definition **name** (`depends_on`,
    `category_member`, a custom name) because the importer prepends the seed
    convention `rel_`; the core speaks record keys. The mapping is read once per
    export from `GET /relation-definitions`; a definition without a name (or an
    unreachable registry) falls back to stripping `rel_` from the key.
    """
    names: dict[str, str] = {}
    try:
        response = await active.get(f"{core_url}/relation-definitions")
        response.raise_for_status()
        for definition in response.json() or []:
            if not isinstance(definition, dict):
                continue
            key = _record_key(definition.get("id"))
            name = str(definition.get("name") or "")
            if key and name:
                names[key] = name
    except Exception as exc:  # noqa: BLE001 - the seed convention is the fallback
        logger.warning("relation definitions unavailable, using the seed convention: %s", exc)
    return names


def _definition_name(definition_id: Any, names: Mapping[str, str]) -> str:
    """The MEF reference `type` of a relation-definition id (see `_load_relation_names`)."""
    key = _record_key(definition_id)
    if key in names:
        return names[key]
    return key[4:] if key.startswith("rel_") else key


async def _load_references(
    active: httpx.AsyncClient,
    core_url: str,
    ids: Iterable[Any],
) -> dict[str, dict[str, Any]]:
    """Every edge touching one of `ids`, keyed by edge id (raw rows).

    `POST /references/query` answers both directions for up to
    [`REFERENCE_QUERY_SIZE`] ids and pages the rows, so a hub node (a category with
    thousands of members) cannot overflow a single response.
    """
    edges: dict[str, dict[str, Any]] = {}
    for chunk in _chunks([str(item) for item in ids], REFERENCE_QUERY_SIZE):
        start = 0
        while True:
            response = await active.post(
                f"{core_url}/references/query",
                json={"ids": chunk, "limit": REFERENCE_QUERY_PAGE, "start": start},
            )
            response.raise_for_status()
            body = response.json() or {}
            rows = [row for row in (body.get("references") or []) if isinstance(row, dict)]
            for row in rows:
                edges[str(row.get("id"))] = row
            if not body.get("truncated"):
                break
            start += REFERENCE_QUERY_PAGE
    return edges


def _mef_references(
    edges: Mapping[str, Mapping[str, Any]],
    id_of: Mapping[str, str],
    names: Mapping[str, str],
) -> tuple[list[dict[str, Any]], int]:
    """MEF reference entries for the edges with **both** endpoints in the document.

    An edge pointing out of the document (a category selection keeps its products,
    not the taxonomy) cannot be imported — the core refuses an unresolvable
    endpoint — so it is left out and **counted** instead of silently vanishing.
    Edges to the universal root are dropped without counting: the root exists in
    every instance and the core attaches a top-level object to it by itself.
    """
    entries: list[dict[str, Any]] = []
    skipped = 0
    for edge_id in sorted(edges):
        edge = edges[edge_id]
        source_key = _record_key(edge.get("in"))
        target_key = _record_key(edge.get("out"))
        if source_key == ROOT_KEY or target_key == ROOT_KEY:
            continue
        source = id_of.get(source_key)
        target = id_of.get(target_key)
        if not source or not target:
            skipped += 1
            continue
        entry: dict[str, Any] = {
            "in": source,
            "out": target,
            "type": _definition_name(edge.get("definition_id"), names),
        }
        attributes = edge.get("attributes")
        if isinstance(attributes, dict) and attributes:
            entry["attributes"] = attributes
        entries.append(entry)
    return entries, skipped


def _mef_releases(obj: Mapping[str, Any]) -> list[dict[str, str]]:
    """The channel releases of an object (`GET /objects?releases=true`).

    A release is channel state in its own table, not an attribute — the MEF object
    carries it as `[{channel, status}]` and the import replays it (`_
    import_releases` in `meadow/worker/imports.py`).
    """
    rows = obj.get("releases")
    if not isinstance(rows, list):
        return []
    return [
        {"channel": str(row.get("channel")), "status": str(row.get("status"))}
        for row in rows
        if isinstance(row, dict) and row.get("channel") and row.get("status")
    ]


async def _mef_objects(
    active: httpx.AsyncClient,
    core_url: str,
    objects: list[dict[str, Any]],
    unit_map: Mapping[str, tuple[str, str]],
) -> list[dict[str, Any]]:
    """The `objects` entries of a MEF document (values, quantified re-emitted)."""
    entries: list[dict[str, Any]] = []
    for obj in objects:
        attributes = obj.get("attributes")
        attributes = dict(attributes) if isinstance(attributes, dict) else {}
        await _emit_quantified(active, core_url, attributes, unit_map)
        entry: dict[str, Any] = {
            "id": obj.get("id"),
            "type": obj.get("type"),
            "attributes": attributes,
            "status": obj.get("status") or "draft",
        }
        releases = _mef_releases(obj)
        if releases:
            entry["releases"] = releases
        entries.append(entry)
    return entries


async def _collect_mef_data(
    active: httpx.AsyncClient,
    core_url: str,
    objects: list[dict[str, Any]],
    *,
    include_ancestors: bool = True,
    with_releases: bool = False,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, int]]:
    """Values **and** relations of the exported objects.

    Returns `(objects, references, counters)`. The objects may have grown by the
    **composition ancestors** of the source (`include_ancestors`), and `references`
    are `{in, out, type, attributes?}` entries — exactly what
    `meadow/worker/imports.py` consumes (composition becomes a re-parent, every
    other relation an edge).

    The ancestor walk exists because a MEF object is always written **with** its
    parent and the import refuses an object whose parent is missing: a document
    that names a parent it does not contain, or omits the relation entirely, cannot
    rebuild the structure tree on a fresh system (`plan/detail/relationen.md`). For
    a complete export the walk stops after the first round — every parent is part
    of the document anyway.
    """
    by_key: dict[str, dict[str, Any]] = {
        _record_key(obj.get("id")): obj for obj in objects if obj.get("id")
    }
    names = await _load_relation_names(active, core_url)
    edges = await _load_references(active, core_url, [str(obj["id"]) for obj in by_key.values()])
    added = 0
    if include_ancestors:
        for _ in range(MAX_ANCESTOR_ROUNDS):
            missing = sorted(
                {
                    _record_key(edge.get("out"))
                    for edge in edges.values()
                    if _record_key(edge.get("definition_id")) == COMPOSITION_DEFINITION
                    and _record_key(edge.get("in")) in by_key
                    and _record_key(edge.get("out")) not in by_key
                    and _record_key(edge.get("out")) != ROOT_KEY
                }
            )
            if not missing:
                break
            fetched = await _load_objects_by_ids(
                active,
                core_url,
                missing,
                extra={"releases": "true"} if with_releases else None,
            )
            if not fetched:
                # A parent that does not exist (deleted). Its edge cannot be part
                # of the document; it is reported as skipped below.
                logger.warning("composition parent(s) missing in the system: %s", missing)
                break
            for obj in fetched:
                key = _record_key(obj.get("id"))
                if key and key not in by_key:
                    by_key[key] = obj
                    added += 1
            edges.update(
                await _load_references(
                    active, core_url, [f"object:{_record_key(obj.get('id'))}" for obj in fetched]
                )
            )
    id_of = {
        _record_key(obj.get("id")): str(obj.get("id")) for obj in by_key.values() if obj.get("id")
    }
    references, skipped = _mef_references(edges, id_of, names)
    counters = {
        "references": len(references),
        "skipped_references": skipped,
        "added_ancestors": added,
    }
    return list(by_key.values()), references, counters


def _mef_document(
    objects: list[dict[str, Any]],
    references: list[dict[str, Any]],
) -> dict[str, Any]:
    """The MEF document envelope (`objects` + `references` + header)."""
    return {
        "format_version": "1.0",
        "source": "meadow",
        "created_at": _now_iso_utc(),
        "objects": objects,
        "references": references,
    }


def build_export_handler(
    core_url: str,
    *,
    api_key: str | None = None,
    client: httpx.AsyncClient | None = None,
) -> JobHandler:
    """Builds the `outbound_export` (json_export) job handler.

    `client` may be injected for tests; otherwise a short-lived client is
    created per job.
    """
    core_url = loopback_host(core_url)
    headers = {"x-api-key": api_key} if api_key else {}

    async def run(active: httpx.AsyncClient, message: JobMessage) -> dict[str, Any]:
        parameters = message.parameters
        # Structured (UI wizard) export: `source` is present.
        if parameters.get("source") is not None:
            return await _structured_export(active, core_url, parameters)

        filter_ = parameters.get("filter") or {}
        if not isinstance(filter_, dict):
            raise RuntimeError("`parameters.filter` must be an object")
        object_type = filter_.get("type")
        status = filter_.get("status")
        limit = int(parameters.get("limit") or DEFAULT_EXPORT_LIMIT)

        query: dict[str, Any] = {"limit": min(max(limit, 1), 1000), "releases": "true"}
        if object_type:
            query["type"] = str(object_type)
        response = await active.get(f"{core_url}/objects", params=query)
        response.raise_for_status()
        objects = response.json()

        if status:
            objects = [obj for obj in objects if obj.get("status") == str(status)]

        objects, references, counters = await _collect_mef_data(
            active, core_url, objects, with_releases=True
        )
        unit_map = await _load_unit_map(active, core_url)

        mef_objects = await _mef_objects(active, core_url, objects, unit_map)

        logger.info(
            "json_export exported %d objects with %d references",
            len(mef_objects),
            len(references),
        )
        return {
            "connector": "json_export",
            "exported": len(mef_objects),
            "document": _mef_document(mef_objects, references),
            **counters,
        }

    async def handle(message: JobMessage) -> dict[str, Any]:
        if client is not None:
            return await run(client, message)
        async with httpx.AsyncClient(timeout=60.0, headers=headers) as active:
            return await run(active, message)

    return handle


__all__ = ["build_export_handler"]
