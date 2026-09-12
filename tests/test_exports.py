"""Tests for the structured export serializer (`meadow/worker/exports.py`).

Covers the `url` type in the export: XLSX cells of a `url` attribute become
clickable hyperlinks, everything else stays plain text.
"""

from __future__ import annotations

import io
from typing import Any

from openpyxl import load_workbook

from meadow.worker.exports import _serialize_xlsx, _url_columns


def _columns() -> list[str]:
    return ["name", "homepage", "homepage_de", "note"]


def _rows() -> list[dict[str, Any]]:
    return [
        {
            "name": "Bench",
            "homepage": "https://example.com/a",
            "homepage_de": "https://example.de/a",
            "note": "kein Link",
        },
        {
            "name": "Chair",
            "homepage": "kaputt",
            "homepage_de": "",
            "note": "https://example.com/nur-text",
        },
    ]


def test_url_columns_cover_multilingual_columns() -> None:
    assert _url_columns(_columns(), {"homepage"}) == {"homepage", "homepage_de"}
    assert _url_columns(_columns(), set()) == set()
    assert _url_columns(_columns(), {"other"}) == set()


def test_xlsx_url_cells_become_hyperlinks() -> None:
    columns = _columns()
    data = _serialize_xlsx(columns, _rows(), _url_columns(columns, {"homepage"}))
    sheet = load_workbook(io.BytesIO(data)).active
    header = [cell.value for cell in sheet[1]]
    assert header == columns
    column = header.index("homepage") + 1

    linked = sheet.cell(row=2, column=column)
    assert linked.value == "https://example.com/a"
    assert linked.hyperlink is not None
    assert linked.hyperlink.target == "https://example.com/a"
    assert linked.font.underline == "single"

    # Per-language columns of the attribute are links as well.
    language = sheet.cell(row=2, column=header.index("homepage_de") + 1)
    assert language.hyperlink is not None
    assert language.hyperlink.target == "https://example.de/a"

    # A value that is not a URL stays plain text (never a dead hyperlink).
    assert sheet.cell(row=3, column=column).hyperlink is None
    assert sheet.cell(row=3, column=header.index("homepage_de") + 1).hyperlink is None

    # Only `url` columns are linked: a `text` attribute holding a URL stays text.
    note = sheet.cell(row=3, column=header.index("note") + 1)
    assert note.value == "https://example.com/nur-text"
    assert note.hyperlink is None


def test_xlsx_without_url_columns_is_unchanged() -> None:
    columns = ["name", "note"]
    rows = [{"name": "Bench", "note": "https://example.com/a"}]
    sheet = load_workbook(io.BytesIO(_serialize_xlsx(columns, rows))).active
    assert [cell.value for cell in sheet[1]] == columns
    assert sheet.cell(row=2, column=2).hyperlink is None
